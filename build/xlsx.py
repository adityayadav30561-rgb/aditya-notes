#!/usr/bin/env python3
"""Build SSB_Notes.xlsx: every sheet and guide in one workbook, one tab each."""
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "SSB_Notes.xlsx")

QUESTION_SHEETS = [("PIQ Questions", "questions.md"), ("Service Knowledge", "service_knowledge.md")]
GUIDE_SHEETS = [("SRT Guide", "srt_guide.md"), ("TAT Guide", "tat_guide.md")]

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
GROUP_FILL = PatternFill("solid", fgColor="D9E2F3")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def plain(text):
    return text.replace("**", "").strip()


def question_rows(path):
    n, sec = 0, ""
    for line in open(path, encoding="utf-8"):
        if line.startswith("## "):
            sec = line[3:].strip()
        elif line.startswith("- "):
            n += 1
            q, _, a = line[2:].partition(" :: ")
            yield [n, sec, q.strip(), a.strip()]


def guide_rows(path):
    sec, group = "", ""
    for line in open(path, encoding="utf-8"):
        line = line.rstrip("\n")
        if line.startswith("## "):
            sec, group = line[3:].strip(), ""
        elif line.startswith("- "):
            body = line[2:]
            m = re.match(r"\*\*(.+?)\*\*\s*(.*)", body)
            if m:
                yield [sec, group, m.group(1).rstrip(":").strip(), plain(m.group(2))]
            else:
                yield [sec, group, "", plain(body)]
        elif line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|", 1)]
            if len(cells) == 1:
                group = plain(cells[0])
            else:
                yield [sec, group, plain(cells[0]), plain(cells[1])]


def write_sheet(ws, headers, widths, rows, title=None, tagline=None):
    r = 1
    if title:
        ws.cell(r, 1, title).font = Font(name=FONT, size=14, bold=True)
        r += 1
    if tagline:
        ws.cell(r, 1, tagline).font = Font(name=FONT, size=10, italic=True)
        r += 1
    if title or tagline:
        r += 1
    head = r
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(r, c, h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center")
        cell.border = BORDER
        ws.column_dimensions[cell.column_letter].width = w
    for row in rows:
        r += 1
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = WRAP
            cell.border = BORDER
    ws.freeze_panes = ws.cell(head + 1, 1)
    ws.auto_filter.ref = f"A{head}:{ws.cell(r, len(headers)).coordinate}"
    ws.print_title_rows = f"{head}:{head}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return head


def shade_groups(ws, head, key_col):
    """Shade the first row of each new section so sections stand out."""
    prev = None
    for r in range(head + 1, ws.max_row + 1):
        val = ws.cell(r, key_col).value
        if val != prev:
            for c in range(1, ws.max_column + 1):
                ws.cell(r, c).fill = GROUP_FILL
            prev = val


def main():
    wb = Workbook()
    wb.remove(wb.active)
    for name, src in QUESTION_SHEETS:
        ws = wb.create_sheet(name)
        rows = list(question_rows(os.path.join(ROOT, src)))
        head = write_sheet(ws, ["#", "Section", "Question", "Answer"], [6, 28, 55, 80], rows)
        shade_groups(ws, head, 2)
        print(f"{name}: {len(rows)} questions")
    for name, src in GUIDE_SHEETS:
        path = os.path.join(ROOT, src)
        lines = open(path, encoding="utf-8").read().splitlines()
        title = next((l[2:] for l in lines if l.startswith("# ")), name)
        tagline = next((plain(l[2:]) for l in lines if l.startswith("> ")), None)
        ws = wb.create_sheet(name)
        rows = list(guide_rows(path))
        head = write_sheet(ws, ["Section", "Group", "Point / Situation", "Detail"], [30, 26, 34, 95],
                           rows, title, tagline)
        shade_groups(ws, head, 1)
        print(f"{name}: {len(rows)} rows")
    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
